"""Exercise real DOCX generation, revisions, comments and render failures."""
import argparse
from contextlib import closing
import hashlib
import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree
from PIL import Image
import pypdfium2 as pdfium


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--skill', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    script = args.skill.resolve() / 'scripts/documents.py'
    args.output.mkdir(parents=True, exist_ok=False)
    root = args.output.resolve()

    def invoke(*arguments, success=True):
        result = subprocess.run([sys.executable, '-B', str(script), *map(str, arguments)],
                                capture_output=True, text=True, cwd=root, timeout=240)
        assert (result.returncode == 0) == success, result.stderr + result.stdout
        return json.loads(result.stdout) if success else result.stderr

    doctor = invoke('doctor')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    for name in ['Normal', 'Title', 'Heading 1']:
        style = doc.styles[name]
        font = 'Noto Sans CJK SC'
        style.font.name = font
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)
    doc.styles['Normal'].font.size = Pt(10)
    doc.add_heading('Documents acceptance', 0)
    doc.add_paragraph('\u6587\u6863\u751f\u6210\u4e0e\u5ba1\u9605\u9a8c\u6536')
    doc.add_paragraph('Status: Draft. Keep this formatting.').runs[0].bold = True
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['Item', 'Owner', 'Result']):
        cell.text = text
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for index in range(80):
        for cell, text in zip(table.add_row().cells, [str(index + 1), '\u6587\u6863\u56e2\u961f', 'Checked']):
            cell.text = text
    source = root / 'source.docx'
    doc.save(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    comment = invoke('comment', source, '--paragraph', 2, '--text', 'Verify delivery status', '--author', 'Acceptance', '--output', root/'commented.docx')
    assert comment['comments'][0]['text'] == 'Verify delivery status'
    with ZipFile(root/'commented.docx') as package:
        xml = etree.fromstring(package.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        assert len(xml.findall('.//w:commentRangeStart', ns)) == 1
        assert len(xml.findall('.//w:commentRangeEnd', ns)) == 1
        assert b'comments.xml' in package.read('word/_rels/document.xml.rels')
    revision = invoke('replace', source, '--old', 'Draft', '--new', 'Approved', '--author', 'Acceptance', '--output', root/'revised.docx')
    assert revision['insertions'] == revision['deletions'] == 1
    with ZipFile(root/'revised.docx') as package:
        xml = etree.fromstring(package.read('word/document.xml'))
        assert xml.find('.//w:del//w:delText', ns).text == 'Draft'
        assert xml.find('.//w:ins//w:t', ns).text == 'Approved'
        assert xml.find('.//w:ins//w:rPr/w:b', ns) is not None
    assert 'new file' in invoke('replace', source, '--old', 'Draft', '--new', 'Approved', '--author', 'Acceptance', '--output', source, success=False)
    assert 'exactly one' in invoke('replace', source, '--old', 'absent', '--new', 'Approved', '--author', 'Acceptance', '--output', root/'invalid.docx', success=False)
    ambiguous = Document()
    ambiguous.add_paragraph('Draft')
    ambiguous.add_paragraph('Draft')
    ambiguous.save(root/'ambiguous.docx')
    assert 'exactly one' in invoke('replace', root/'ambiguous.docx', '--old', 'Draft', '--new', 'Approved', '--author', 'Acceptance', '--output', root/'ambiguous-out.docx', success=False)
    split = Document()
    paragraph = split.add_paragraph()
    paragraph.add_run('Dr').bold = True
    paragraph.add_run('aft')
    split.save(root/'split.docx')
    assert 'exactly one' in invoke('replace', root/'split.docx', '--old', 'Draft', '--new', 'Approved', '--author', 'Acceptance', '--output', root/'split-out.docx', success=False)
    receipt = invoke('render', root/'commented.docx', '--output', root/'qa')
    assert receipt['pageCount'] >= 2
    assert receipt['sha256'] == hashlib.sha256((root/'commented.docx').read_bytes()).hexdigest()
    for page in receipt['pages']:
        with Image.open(root/page['path']) as image:
            assert image.width > 800 and image.height > 1000
            assert image.convert('L').getextrema()[0] < 240
    with closing(pdfium.PdfDocument(str(root/'qa/document.pdf'))) as pdf:
        text = ''
        for page in pdf:
            with closing(page.get_textpage()) as words:
                text += words.get_text_bounded()
            page.close()
        assert '80' in text and '\u6587\u6863' in text
    assert 'new directory' in invoke('render', source, '--output', root/'qa', success=False)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    summary = {'status': 'passed', 'pageCount': receipt['pageCount'], 'python': doctor['python'],
               'soffice': doctor['soffice'], 'checks': ['comments-and-anchors', 'tracked-text-and-style',
               'source-preserved', 'ambiguous-edit-rejected', 'cross-run-edit-rejected', 'fresh-render', 'CJK-text', 'PNG-pages'],
               'visualReview': 'pending'}
    (root/'smoke.json').write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary))


if __name__ == '__main__':
    main()
