"""Real editable PPTX fixture, conservative-edit integrity and rendering acceptance."""
import argparse
from contextlib import closing
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import zipfile

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
import pypdfium2 as pdfium


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--skill', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    script = args.skill.resolve() / 'scripts/presentations.py'
    args.output.mkdir(parents=True, exist_ok=False)
    root = args.output.resolve()

    def invoke(*arguments, success=True):
        result = subprocess.run([sys.executable, '-B', str(script), *map(str, arguments)],
                                cwd=root, text=True, capture_output=True, timeout=240)
        assert (result.returncode == 0) == success, result.stderr + result.stdout
        return json.loads(result.stdout) if success else result.stderr

    doctor = invoke('doctor')
    with Image.new('RGB', (800, 450), '#2160CB') as picture:
        drawing = ImageDraw.Draw(picture)
        drawing.rounded_rectangle((70, 70, 730, 380), radius=24, fill='#42C9B0')
        drawing.line([(130, 300), (300, 230), (470, 270), (660, 120)], fill='white', width=12)
        picture.save(root/'trend.png')
    chinese = '\u4e2d\u6587\u6f14\u793a\u9a8c\u6536'
    rows = [['Workstream', 'Owner', 'Status'], ['Generation', 'Team A', 'Ready'], ['Review', 'Team B', 'Ready']]
    spec = {'title': chinese, 'theme': 'light', 'slides': [
        {'title': chinese, 'layout': 'cover', 'notes': 'Opening notes and source: https://example.com', 'elements': [
            {'type': 'text', 'x': 0.9, 'y': 3.5, 'w': 11, 'h': 1, 'text': 'Editable slides - acceptance fixture', 'fontSize': 26}]},
        {'title': 'Native table and image', 'notes': 'Two body rows plus one header.', 'elements': [
            {'type': 'table', 'x': 0.8, 'y': 1.8, 'w': 7, 'h': 3, 'rows': rows, 'fontSize': 18},
            {'type': 'image', 'x': 8.3, 'y': 2, 'w': 4.1, 'h': 3, 'path': 'trend.png', 'altText': 'Illustrative trend'}]},
        {'title': 'Native chart data', 'elements': [
            {'type': 'chart', 'x': 0.9, 'y': 1.6, 'w': 11.5, 'h': 4.8, 'chartType': 'bar',
             'series': [{'name': 'Complete', 'labels': ['Q1', 'Q2', 'Q3'], 'values': [20, 35, 50]}]}]},
        {'title': 'Next steps', 'elements': [
            {'type': 'shape', 'shape': 'roundRect', 'x': 0.8, 'y': 1.8, 'w': 11.6, 'h': 3.5},
            {'type': 'text', 'x': 1.2, 'y': 2.2, 'w': 10.8, 'h': 2.5,
             'text': 'Inspect editable objects\nReview all rendered slides\nSave and inspect the latest file', 'fontSize': 24}]}]}
    spec_path = root/'deck.json'
    spec_path.write_text(json.dumps(spec), encoding='utf-8')
    source = root/'source.pptx'
    info = invoke('build', spec_path, '--output', source)
    assert info['slideCount'] == 4 and not info['warnings']
    for page in info['slides']:
        ids = [shape['shapeId'] for shape in page['shapes']]
        assert len(ids) == len(set(ids)), 'Native objects must have unique shape IDs'
    assert chinese in info['slides'][0]['allText'] and 'Opening notes' in info['slides'][0]['notes']
    table = next(shape for shape in info['slides'][1]['shapes'] if shape['type'] == 'table')
    assert table['rows'] == rows
    assert any(shape['type'] == 'image' for shape in info['slides'][1]['shapes'])
    chart = next(shape for shape in info['slides'][2]['shapes'] if shape['type'] == 'chart')
    assert chart['editableWorkbook'] and chart['series'][0]['values'] == [20, 35, 50]
    assert chart['categories'] == ['Q1', 'Q2', 'Q3']
    title = next(shape for shape in info['slides'][0]['shapes'] if shape.get('text') == chinese)
    old_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    updated = root/'updated.pptx'
    edit_args = ['replace-text', source, '--slide', '1', '--shape-id', str(title['shapeId']),
                 '--old', chinese, '--new', 'Updated title', '--expected-sha256', old_hash]
    invoke(*edit_args, '--output', updated)
    assert 'Updated title' in invoke('inspect', updated)['slides'][0]['allText']
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(updated) as after:
        assert before.namelist() == after.namelist()
        changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        assert changed == ['ppt/slides/slide1.xml'], changed
    assert 'new PPTX' in invoke(*edit_args, '--output', source, success=False)
    stale = edit_args.copy(); stale[-1] = '0' * 64
    assert 'Source changed' in invoke(*stale, '--output', root/'stale.pptx', success=False)
    cross = Presentation(source)
    paragraph = cross.slides[0].shapes[1].text_frame.paragraphs[0]
    paragraph.clear(); paragraph.add_run().text = 'Split'; paragraph.add_run().text = ' title'
    cross.save(root/'cross-run.pptx')
    cross_info = invoke('inspect', root/'cross-run.pptx')
    assert 'single run' in invoke('replace-text', root/'cross-run.pptx', '--output', root/'cross-edited.pptx',
        '--slide', '1', '--shape-id', str(title['shapeId']), '--old', 'Split title', '--new', 'After',
        '--expected-sha256', cross_info['sha256'], success=False)
    bad_spec = json.loads(json.dumps(spec)); bad_spec['slides'][0]['elements'][0]['w'] = 40
    (root/'invalid.json').write_text(json.dumps(bad_spec))
    invoke('build', root/'invalid.json', '--output', root/'invalid.pptx', success=False)
    assert not (root/'invalid.pptx').exists()
    # Presentation order, rather than numeric slide filenames, determines reading and editing.
    reordered = Presentation(source)
    ids = reordered.slides._sldIdLst
    ids.insert(0, ids[-1]); reordered.save(root/'reordered.pptx')
    assert 'Next steps' in invoke('inspect', root/'reordered.pptx')['slides'][0]['allText']
    # Real relationship rejection (not file-name heuristics).
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(root/'external.pptx', 'w') as after:
        for part in before.infolist():
            data = before.read(part.filename)
            if part.filename == 'ppt/slides/_rels/slide2.xml.rels':
                tree = etree.fromstring(data)
                image_rel = next(rel for rel in tree if rel.get('Type').endswith('/image'))
                image_rel.set('TargetMode', 'External'); image_rel.set('Target', 'https://example.com/image.png')
                data = etree.tostring(tree)
            after.writestr(part, data)
    assert 'external' in invoke('render', root/'external.pptx', '--output', root/'external-qa', success=False)
    receipts = []
    for name in ['source', 'updated']:
        receipt = invoke('render', root/(name+'.pptx'), '--output', root/(name+'-qa'))
        assert receipt['slideCount'] == receipt['pageCount'] == 4
        for page in receipt['pages']:
            with Image.open(root/page['path']) as image:
                assert image.width > 1000 and image.convert('L').getextrema()[0] < 240
        receipts.append(receipt)
    variants = {'title': 'Chart variants', 'theme': 'navy', 'slides': [
        {'title': 'Line chart', 'elements': [{'type': 'chart', 'chartType': 'line',
            'x': 0.9, 'y': 1.6, 'w': 11.5, 'h': 4.8, 'series': [
                {'name': 'Plan', 'labels': ['Q1', 'Q2', 'Q3'], 'values': [15, 30, 45]},
                {'name': 'Actual', 'labels': ['Q1', 'Q2', 'Q3'], 'values': [20, 35, 50]}]}]},
        {'title': 'Pie chart', 'elements': [{'type': 'chart', 'chartType': 'pie',
            'x': 0.9, 'y': 1.6, 'w': 11.5, 'h': 4.8, 'series': [
                {'name': 'Share', 'labels': ['A', 'B', 'C'], 'values': [20, 35, 45]}]}]}]}
    (root/'variants.json').write_text(json.dumps(variants))
    variant_info = invoke('build', root/'variants.json', '--output', root/'variants.pptx')
    for page in variant_info['slides']:
        assert next(shape for shape in page['shapes'] if shape['type'] == 'chart')['editableWorkbook']
    receipts.append(invoke('render', root/'variants.pptx', '--output', root/'variants-qa'))
    assert 'new directory' in invoke('render', source, '--output', root/'source-qa', success=False)
    with closing(pdfium.PdfDocument(str(root/'source-qa/presentation.pdf'))) as pdf:
        with closing(pdf[0]) as page:
            with closing(page.get_textpage()) as text:
                assert chinese in text.get_text_bounded()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == old_hash
    result = {'status': 'passed', 'checks': ['native-text-table-image-chart', 'embedded-chart-workbook',
        'unique-shape-IDs', 'line-pie-and-navy-theme', 'speaker-notes', 'CJK-rendered-text', 'slide-order', 'all-slides-rendered', 'source-preserved',
        'unchanged-other-parts', 'stale-and-cross-run-rejected', 'external-media-rejected', 'fresh-QA'],
        'runtime': doctor, 'renders': receipts, 'visualReview': 'pending'}
    (root/'smoke.json').write_text(json.dumps(result, indent=2))
    print(json.dumps({key: value for key, value in result.items() if key != 'renders'}))


if __name__ == '__main__':
    main()
