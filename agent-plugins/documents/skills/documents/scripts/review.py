"""Conservative review helpers: preserve original packages and reject complex edits."""
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def inspect_document(source):
    doc = Document(source)
    with ZipFile(source) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'), etree.XMLParser(resolve_entities=False, no_network=True))
    return {
        'paragraphs': [{'index': i, 'text': p.text} for i, p in enumerate(doc.paragraphs)],
        'tables': [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables],
        'comments': [{'id': c.comment_id, 'author': c.author, 'text': c.text} for c in doc.comments],
        'insertions': len(xml.findall('.//w:ins', NS)),
        'deletions': len(xml.findall('.//w:del', NS))}


def tracked_replace(source, output, old, new, author):
    if not old or '\n' in old or '\n' in new or '\t' in new:
        raise ValueError('Replacement requires non-empty, single-line ordinary text')
    with ZipFile(source) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'), etree.XMLParser(resolve_entities=False, no_network=True))
        if xml.findall('.//w:ins', NS) or xml.findall('.//w:del', NS):
            raise ValueError('Existing revisions require explicit review before adding replacements')
        # A match split over multiple runs is deliberately not flattened.
        candidates = [r for r in xml.findall('.//w:r', NS)
                      if old in ''.join(r.itertext()) and r.find('w:t', NS) is not None]
        if len(candidates) != 1:
            raise ValueError('Expected exactly one ordinary run match; inspect ambiguous or cross-run text manually')
        run = candidates[0]
        children = list(run)
        texts = run.findall('w:t', NS)
        if len(texts) != 1 or any(child.tag not in [qn('w:rPr'), qn('w:t')] for child in children):
            raise ValueError('Target run contains unsupported fields or objects')
        if run.getparent().tag != qn('w:p'):
            raise ValueError('Target is inside a hyperlink or other structured container')
        value = texts[0].text or ''
        if value.count(old) != 1:
            raise ValueError('Target text is ambiguous')
        before, after = value.split(old)
        parent, position = run.getparent(), run.getparent().index(run)
        stamp = datetime.now(timezone.utc).isoformat(timespec='seconds').replace('+00:00', 'Z')
        ids = [int(x) for x in xml.xpath('//@w:id', namespaces=NS) if str(x).isdigit()]
        next_id = max(ids, default=0) + 1

        def text_run(text, deleted=False):
            result = OxmlElement('w:r')
            properties = run.find(qn('w:rPr'))
            if properties is not None:
                result.append(deepcopy(properties))
            element = OxmlElement('w:delText' if deleted else 'w:t')
            element.set(qn('xml:space'), 'preserve')
            element.text = text
            result.append(element)
            return result

        replacements = [text_run(before)] if before else []
        for tag, text in [('del', old), ('ins', new)]:
            if not text:
                continue
            revision = OxmlElement('w:' + tag)
            for key, val in [('id', str(next_id)), ('author', author), ('date', stamp)]:
                revision.set(qn('w:' + key), val)
            next_id += 1
            revision.append(text_run(text, tag == 'del'))
            replacements.append(revision)
        if after:
            replacements.append(text_run(after))
        parent.remove(run)
        for offset, element in enumerate(replacements):
            parent.insert(position + offset, element)
        with ZipFile(output, 'w') as dest:
            for entry in archive.infolist():
                content = etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True) if entry.filename == 'word/document.xml' else archive.read(entry)
                dest.writestr(entry, content)


def operate(args):
    source = Path(args.input).resolve()
    if args.action == 'inspect':
        return inspect_document(source)
    output = Path(args.output).resolve()
    if source == output or output.exists():
        raise ValueError('Output must be a new file; the original is preserved')
    output.parent.mkdir(parents=True, exist_ok=True)
    if not args.author.strip():
        raise ValueError('A review author is required')
    if args.action == 'replace':
        tracked_replace(source, output, args.old, args.new, args.author)
    else:
        doc = Document(source)
        if args.paragraph < 0 or args.paragraph >= len(doc.paragraphs):
            raise ValueError('Paragraph index is out of range')
        paragraph = doc.paragraphs[args.paragraph]
        if not paragraph.runs or not paragraph.text.strip():
            raise ValueError('Comment target must contain text')
        if paragraph._p.xpath('.//w:fldChar | .//w:instrText | .//w:hyperlink | .//w:ins | .//w:del | .//w:commentRangeStart | .//w:commentRangeEnd'):
            raise ValueError('Complex paragraph requires explicit comment anchoring')
        doc.add_comment(paragraph.runs, text=args.text, author=args.author)
        doc.save(output)
    return {'output': str(output), **inspect_document(output)}
