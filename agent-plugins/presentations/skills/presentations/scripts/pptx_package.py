"""Bounded OOXML inspection and exact, single-run text edits without roundtripping."""
import hashlib
import io
import json
from pathlib import Path, PurePosixPath
import posixpath
from urllib.parse import unquote
import zipfile

from lxml import etree
from pptx import Presentation

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'rel': 'http://schemas.openxmlformats.org/package/2006/relationships'}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def xml(data):
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    tree = etree.fromstring(data, parser)
    if tree.getroottree().docinfo.doctype:
        raise ValueError('DTD declarations are unsupported')
    return tree


def local_path(path):
    try:
        return Path(path).resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return str(Path(path).resolve())


def new_output(source, output):
    output = Path(output).resolve()
    if output.suffix.lower() != '.pptx' or output.exists() or (source and output == Path(source).resolve()):
        raise ValueError('Output must be a new PPTX path; preserve the original')
    output.parent.mkdir(parents=True, exist_ok=True)
    return output


class Package:
    def __init__(self, source):
        self.source = Path(source).resolve()
        if self.source.suffix.lower() != '.pptx' or not self.source.is_file():
            raise ValueError('Input must be an existing PPTX file')
        if self.source.stat().st_size > 100 * 1024 * 1024:
            raise ValueError('PPTX exceeds 100 MiB')
        self.data = self.source.read_bytes()
        self.zip = zipfile.ZipFile(io.BytesIO(self.data))
        infos = self.zip.infolist()
        self.names = set(self.zip.namelist())
        if len(infos) > 10000 or len(infos) != len(self.names) or sum(i.file_size for i in infos) > 300 * 1024 * 1024:
            raise ValueError('Duplicate ZIP entries or package size/count limit exceeded')
        for info in infos:
            path = PurePosixPath(info.filename)
            if path.is_absolute() or '..' in path.parts or '\\' in info.filename or info.flag_bits & 1:
                raise ValueError('Unsupported ZIP entry or encrypted archive')
            if info.file_size > 30 * 1024 * 1024:
                raise ValueError('Package part exceeds 30 MiB')
            if 'vbaproject' in info.filename.lower() or info.filename.startswith('_xmlsignatures/'):
                raise ValueError('Macro-enabled or signed decks are unsupported')
        bad = self.zip.testzip()
        if bad:
            raise ValueError('ZIP integrity check failed')
        self.external = []
        self.relationships = {}
        for name in sorted(self.names):
            if not name.endswith('.rels'):
                continue
            directory = posixpath.dirname(posixpath.dirname(name))
            relations = {}
            for rel in xml(self.zip.read(name)):
                rid, target, kind = rel.get('Id'), rel.get('Target'), rel.get('Type', '')
                if not rid or not target or rid in relations:
                    raise ValueError('Invalid or duplicate relationship')
                if rel.get('TargetMode') == 'External':
                    self.external.append({'part': name, 'type': kind.rsplit('/', 1)[-1]})
                    relations[rid] = (None, kind)
                    continue
                resolved = posixpath.normpath(posixpath.join(directory, unquote(target)).lstrip('/'))
                if resolved not in self.names:
                    raise ValueError('Broken package relationship: ' + resolved)
                relations[rid] = (resolved, kind)
            self.relationships[name] = relations
        root = xml(self.zip.read('ppt/presentation.xml'))
        rels = self.relationships['ppt/_rels/presentation.xml.rels']
        self.slides = []
        for slide in root.findall('p:sldIdLst/p:sldId', NS):
            target, kind = rels[slide.get('{%s}id' % NS['r'])]
            if not target or not kind.endswith('/slide'):
                raise ValueError('Invalid slide relationship')
            self.slides.append(target)
        if not 1 <= len(self.slides) <= 100 or len(set(self.slides)) != len(self.slides):
            raise ValueError('Expected 1..100 distinct slides')
        for part in self.slides:
            ids = self.tree(part).xpath('.//p:cNvPr/@id', namespaces=NS)
            if len(ids) != len(set(ids)):
                raise ValueError('Duplicate shape IDs in ' + part)

    def tree(self, part):
        return xml(self.zip.read(part))

    def assert_renderable(self):
        if any(item['type'] != 'hyperlink' for item in self.external):
            raise ValueError('Rendering external media/data relationships is unsupported')
        for part in self.slides:
            tree = self.tree(part)
            if tree.get('show') == '0':
                raise ValueError('Hidden slides require explicit normalization before rendering in v1')
            if tree.findall('.//p:oleObj', NS):
                raise ValueError('Rendering embedded OLE objects is unsupported in v1')


def normalize_generated_ids(path):
    # PptxGenJS 4.0.1 uses a separate counter for table IDs, which can collide
    # with titles. Our generated decks have no animations/connectors referencing
    # shape IDs; assign unique IDs before publishing. Never normalize uploads.
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, 'w') as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith('ppt/slides/') and info.filename.endswith('.xml'):
                tree = xml(data)
                for index, props in enumerate(tree.xpath('.//p:cNvPr', namespaces=NS), start=1):
                    props.set('id', str(index))
                data = etree.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
            target.writestr(info, data)
    Path(path).write_bytes(buffer.getvalue())


def inspect_deck(source):
    package = Package(source)
    deck = Presentation(io.BytesIO(package.data))
    pages = []
    warnings = []
    for index, slide in enumerate(deck.slides):
        shapes = []
        for shape in slide.shapes:
            item = {'shapeId': shape.shape_id, 'name': shape.name, 'type': 'other',
                    'bounds': [round(getattr(shape, name) / 914400, 3) for name in ['left', 'top', 'width', 'height']]}
            if shape.has_text_frame:
                item.update(type='text', text=shape.text,
                            runs=[run.text for para in shape.text_frame.paragraphs for run in para.runs])
            elif shape.has_table:
                item.update(type='table', rows=[[cell.text for cell in row.cells] for row in shape.table.rows])
            elif shape.has_chart:
                item.update(type='chart', chartType=str(shape.chart.chart_type),
                            series=[{'name': series.name, 'values': list(series.values)} for series in shape.chart.series])
                item['categories'] = [str(category.label) for category in shape.chart.plots[0].categories]
                item['editableWorkbook'] = bool(shape.chart.part.chart_workbook.xlsx_part)
            elif shape.shape_type == 13:
                item['type'] = 'image'
            # Grouped/master/SmartArt content is not a v1 editing target.
            if shape.left < -9144 or shape.top < -9144 or shape.left + shape.width > deck.slide_width + 9144 or shape.top + shape.height > deck.slide_height + 9144:
                warnings.append({'slide': index + 1, 'shapeId': shape.shape_id, 'issue': 'outside-slide-bounds'})
            shapes.append(item)
        tree = package.tree(package.slides[index])
        pages.append({'slide': index + 1, 'hidden': tree.get('show') == '0', 'shapes': shapes,
                      'allText': '\n'.join(tree.xpath('.//a:t/text()', namespaces=NS)),
                      'notes': slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ''})
    return {'input': local_path(source), 'sha256': sha(package.data), 'slideCount': len(pages),
            'width': round(deck.slide_width / 914400, 3), 'height': round(deck.slide_height / 914400, 3),
            'slides': pages, 'warnings': warnings, 'externalRelationships': package.external}


def replace_text(source, output, slide, shape_id, old, new, expected_sha):
    package = Package(source)
    output = new_output(source, output)
    if expected_sha and sha(package.data) != expected_sha:
        raise ValueError('Source changed since inspection; inspect the latest saved file again')
    if not 1 <= slide <= len(package.slides) or not old or old == new:
        raise ValueError('Specify a valid slide and a nonempty, changed exact text')
    if len(new) > 4000 or any(ord(char) < 32 and char not in '\t\n\r' for char in new):
        raise ValueError('Replacement is too large or contains invalid XML characters')
    part = package.slides[slide - 1]
    tree = package.tree(part)
    targets = tree.xpath('./p:cSld/p:spTree/p:sp[p:nvSpPr/p:cNvPr/@id=$id]', namespaces=NS, id=str(shape_id))
    if len(targets) != 1:
        raise ValueError('Editing supports one top-level text shape; tables, charts, groups and masters are unsupported')
    runs = targets[0].xpath('./p:txBody/a:p/a:r/a:t', namespaces=NS)
    matches = [run for run in runs if old in (run.text or '')]
    if len(matches) != 1 or matches[0].text.count(old) != 1:
        raise ValueError('Text must match exactly once inside a single run; cross-run and ambiguous edits are unsupported')
    matches[0].text = matches[0].text.replace(old, new)
    changed = etree.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as archive:
        for info in package.zip.infolist():
            archive.writestr(info, changed if info.filename == part else package.zip.read(info.filename))
    if sha(Path(source).read_bytes()) != sha(package.data):
        raise ValueError('Source changed during editing; retry after saving')
    with output.open('xb') as file:
        file.write(buffer.getvalue())
    try:
        result = inspect_deck(output)
    except Exception:
        output.unlink()
        raise
    return {'output': local_path(output), 'sourceSha256': sha(package.data), 'sha256': result['sha256'],
            'slideCount': result['slideCount'], 'changedPart': part,
            'preservedOtherParts': len(package.names) - 1, 'renderRequired': True}
